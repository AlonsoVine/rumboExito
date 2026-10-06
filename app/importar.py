# -*- coding: utf-8 -*-
"""
importar.py  ·  Trae movimientos de golpe
=========================================
Tres caminos, todos con vista previa antes de guardar nada:
  1. El CSV «Plusvalías y minusvalías» de cada fondo de MyInvestor.
  2. La plantilla genérica (Excel o CSV), rellena a mano o por una IA.
  3. El texto que devuelve la IA, pegado tal cual.

El proceso tiene dos pasos: preparar() lo lee, reconoce los productos y deja un
plan; aplicar() ejecuta ese plan sobre la cartera. La vista previa es aplicar()
sobre una copia, así que lo que ves es exactamente lo que se guardará.
"""

import copy
import csv
import datetime as dt
import io
import re
import unicodedata

from . import almacen, buscar, motor
from .motor import num_es, valor_en

COLUMNAS = ["fecha", "identificador", "nombre", "tipo_producto", "tipo_movimiento",
            "unidades", "importe", "moneda", "comision", "nota"]
COLUMNAS_BANCO = ["fecha", "concepto", "importe", "tipo", "categoria", "titular"]

# Nombres alternativos que se aceptan en la cabecera.
SINONIMOS = {
    "fecha": ["fecha", "date", "fecha operacion", "fecha valor", "fecha ejecucion"],
    "identificador": ["identificador", "isin", "ticker", "simbolo", "codigo", "id"],
    "nombre": ["nombre", "producto", "descripcion", "name"],
    "tipo_producto": ["tipo producto", "tipo_producto", "clase", "tipo de producto"],
    "tipo_movimiento": ["tipo movimiento", "tipo_movimiento", "movimiento", "operacion", "tipo"],
    "unidades": ["unidades", "participaciones", "titulos", "acciones", "cantidad", "units"],
    "importe": ["importe", "importe eur", "total", "amount", "importe total"],
    "moneda": ["moneda", "divisa", "currency"],
    "comision": ["comision", "comisiones", "fee", "gastos"],
    "nota": ["nota", "notas", "comentario", "observaciones"],
}
TIPO_MOV = {
    "compra": "compra", "suscripcion": "compra", "buy": "compra", "aportacion": "compra",
    "venta": "venta", "reembolso": "venta", "sell": "venta",
    "dividendo": "dividendo", "cupon": "dividendo", "dividend": "dividendo", "interes": "dividendo",
    "comision": "comision", "fee": "comision", "custodia": "comision",
    "saldo": "saldo", "valor": "saldo", "valoracion": "saldo",
}
TIPO_PROD = {
    "fondo": "fondo", "fondo de inversion": "fondo", "etf": "etf", "etp": "etf",
    "accion": "accion", "acciones": "accion", "cripto": "cripto", "criptomoneda": "cripto",
    "commodity": "commodity", "materia prima": "commodity", "oro": "commodity",
    "bono": "bono", "renta fija": "bono", "pension": "pension", "plan de pensiones": "pension",
    "efectivo": "efectivo", "cuenta": "efectivo", "cuenta corriente": "efectivo",
    "inmueble": "inmueble", "piso": "inmueble", "vivienda": "inmueble",
    "deuda": "deuda", "prestamo": "deuda", "hipoteca": "deuda", "otro": "otro",
}


def sin_tildes(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t.replace("_", " ")).strip().lower()


def decodifica(crudo):
    """Bytes de un archivo de texto -> str, probando las codificaciones habituales."""
    if isinstance(crudo, str):
        return crudo
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return crudo.decode(enc)
        except UnicodeDecodeError:
            continue
    return crudo.decode("utf-8", "replace")


# ---------------------------------------------------------------- MyInvestor

CAB_MYINVESTOR = {"fecha": ["fecha fiscal", "fecha"],
                  "coste": ["inversion", "inversión", "coste"],
                  "valor": ["valor de mercado", "valor mercado", "valor"]}


def leer_csv_myinvestor(texto):
    """
    Lee el CSV "Plusvalías y minusvalías" de un fondo de MyInvestor.
    Devuelve (lotes, reembolsos): lotes = [[fecha, coste, valor]] de lo que sigues
    teniendo; reembolsos = [[fecha, resultado]] de lo ya vendido, del que el extracto
    solo da la plusvalía.
    """
    texto = decodifica(texto)
    delim = ";" if texto.count(";") >= texto.count(",") else ","
    filas = [f for f in csv.reader(io.StringIO(texto), delimiter=delim)
             if any((c or "").strip() for c in f)]
    if not filas:
        return [], []

    cab = [(c or "").strip().lower() for c in filas[0]]

    def col(clave, defecto):
        for i, c in enumerate(cab):
            if any(a in c for a in CAB_MYINVESTOR[clave]):
                return i
        return defecto

    ic, ico, iv = col("fecha", 0), col("coste", 1), col("valor", 2)
    lotes, reembolsos = [], []
    for fila in filas[1:]:
        if len(fila) <= max(ic, ico, iv):
            continue
        m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", fila[ic])
        if not m:
            continue
        fecha = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
        coste, valor = num_es(fila[ico]), num_es(fila[iv])
        if coste <= 0 and valor <= 0:
            # Participaciones ya reembolsadas: solo queda el resultado fiscal.
            if len(fila) > 3 and num_es(fila[3]):
                reembolsos.append([fecha, round(num_es(fila[3]), 2)])
            continue
        lotes.append([fecha, round(coste, 2), round(valor, 2)])
    lotes.sort(key=lambda x: x[0])
    return lotes, reembolsos


def myinvestor_a_movimientos(producto_id, lotes, reembolsos, serie_vl):
    """
    Convierte los lotes en compras. El extracto no trae participaciones, así que
    se calculan como coste / valor liquidativo del día de compra. No se usa el
    valor de mercado del extracto porque MyInvestor lo calcula con un VL de uno o
    dos días antes y saldrían participaciones erróneas.
    """
    # VL con el que MyInvestor valoró el extracto, para los lotes sin VL de compra.
    estimados = sorted(serie_vl[f] * v / c for f, c, v in lotes if serie_vl.get(f) and c > 0)
    vl_extracto = estimados[len(estimados) // 2] if estimados else None
    movs, sin_vl = [], 0
    for fecha, coste, valor in lotes:
        vc = serie_vl.get(fecha)
        if vc:
            unidades = coste / vc
        elif vl_extracto:
            unidades, sin_vl = valor / vl_extracto, sin_vl + 1
        else:
            unidades, sin_vl = 0.0, sin_vl + 1
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "compra",
                     "unidades": round(unidades, 6), "importe": coste, "nota": "MyInvestor"})
    for fecha, resultado in reembolsos:
        # Del reembolso solo se conoce la plusvalía: se anota como una venta de 0
        # participaciones que cobra ese resultado.
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "venta",
                     "unidades": 0, "importe": resultado,
                     "nota": "Plusvalía de un reembolso (MyInvestor no da la fecha de venta)"})
    return movs, sin_vl


# ---------------------------------------------------------------- lectura de tablas

def leer_tabla(nombre, contenido):
    """Excel (.xlsx), CSV o texto pegado -> [(número de fila, {columna: valor})]."""
    if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        hoja = libro["Movimientos"] if "Movimientos" in libro.sheetnames else libro.worksheets[0]
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    else:
        texto = decodifica(contenido)
        # Lo que devuelve una IA suele venir dentro de un bloque ``` ... ```.
        texto = re.sub(r"^\s*```[a-zA-Z]*\s*$", "", texto, flags=re.M).strip()
        lineas = [l for l in texto.splitlines() if l.strip()]
        if not lineas:
            return [], "El texto está vacío."
        cab = lineas[0]
        delim = max([";", "\t", ","], key=cab.count)
        filas = list(csv.reader(io.StringIO("\n".join(lineas)), delimiter=delim))
    # Se guarda el número de fila real (el que ves en Excel) aunque haya filas en blanco.
    filas = [(n, f) for n, f in enumerate(filas, start=1)
             if f and any(c not in (None, "") and str(c).strip() for c in f)]
    if not filas:
        return [], "No hay ninguna fila con datos."

    cab = [sin_tildes(c) for c in filas[0][1]]
    idx = {}
    for col, alias in SINONIMOS.items():
        for i, c in enumerate(cab):
            if c in alias and i not in idx.values():
                idx[col] = i
                break
    faltan = [c for c in ("fecha", "tipo_movimiento", "importe") if c not in idx]
    if faltan:
        return [], ("No encuentro las columnas " + ", ".join(faltan) + ". La primera fila tiene que ser la "
                    "cabecera de la plantilla: " + ";".join(COLUMNAS))
    if len(filas) < 2:
        if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
            return [], ("La hoja «Movimientos» está vacía: solo tiene la cabecera. Las filas de la hoja "
                        "«Ejemplo» son de muestra y no se importan; escribe tus operaciones en «Movimientos», "
                        "debajo de la cabecera, guarda el archivo y vuelve a subirlo.")
        return [], "Solo hay cabecera: escribe tus operaciones debajo, una por fila."
    salida = []
    for n, f in filas[1:]:
        salida.append((n, {col: (f[i] if i < len(f) else None) for col, i in idx.items()}))
    return salida, None


def lee_fecha(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    t = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        a, me, di = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", t)
        if not m:
            return None
        a = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
        me, di = int(m.group(2)), int(m.group(1))
    try:
        return dt.date(a, me, di).isoformat()
    except ValueError:   # un 31 de febrero, un mes 13...
        return None


def lee_numero(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return abs(float(v))
    t = str(v).strip().replace("€", "").replace("$", "").replace(" ", "")
    if not re.fullmatch(r"[-+]?[\d.,]+", t):
        return "error"
    return abs(num_es(t))


# ---------------------------------------------------------------- preparar el plan

class Plan:
    """Lo que se va a importar: productos nuevos, movimientos, saldos y errores."""

    def __init__(self):
        self.productos_nuevos = []   # [{"ref": "nuevo:1", "datos": {...}, "precio": ..., "fecha": ...}]
        self.movimientos = []        # [{"fila", "producto" (id o ref), campos del movimiento, "marcas"}]
        self.valoraciones = []
        self.flujos = []             # [{"fila", "fecha", "tipo", "importe", "categoria", "titular", "concepto"}]
        self.reemplazar = []         # [(producto id o ref, origen)]: se borran sus importados previos
        self.errores = []            # [{"fila", "mensaje"}]
        self.avisos = []
        self.tipoImport = "activos"  # "activos" o "banco"

    def error(self, fila, mensaje):
        self.errores.append({"fila": fila, "mensaje": mensaje})


def _producto_existente(cfg, identificador, nombre):
    ident = (identificador or "").strip().upper()
    for p in cfg.get("productos", []):
        if ident and ident in ((p.get("identificador") or "").upper(), (p.get("codigo") or "").upper()):
            return p
    nom = sin_tildes(nombre)
    if nom and not ident:
        for p in cfg.get("productos", []):
            if nom in (sin_tildes(p.get("nombre")), sin_tildes(p.get("corto"))):
                return p
    return None


class Resolutor:
    """Encuentra (o propone crear) el producto de cada fila, una sola vez por producto."""

    def __init__(self, cfg, plan, carpeta):
        self.cfg, self.plan, self.carpeta = cfg, plan, carpeta
        self.cache, self.series, self.frescas = {}, {}, set()

    def producto(self, fila, identificador, nombre, tipo_prod):
        clave = ((identificador or "").strip().upper(), sin_tildes(nombre) if not identificador else "")
        if clave in self.cache:
            return self.cache[clave]
        ref = None
        existente = _producto_existente(self.cfg, identificador, nombre)
        if existente:
            ref = existente["id"]
        elif identificador:
            res = buscar.buscar(identificador)
            if res:
                r = res[0]
                fi = r.get("ficha") or {}
                datos = {"nombre": nombre or r.get("nombre") or r["codigo"], "tipo": r["tipo"],
                         "identificador": identificador.strip(), "fuente": r["fuente"], "codigo": r["codigo"],
                         "moneda": r.get("moneda") or "EUR", "vivo": r.get("vivo") or "",
                         "ter": fi.get("ter"), "riesgo": fi.get("riesgo"), "clase": fi.get("clase") or "",
                         "gestora": fi.get("gestora") or "", "largoPlazo": True}
                ref = self._nuevo(datos, r)
            else:
                self.plan.error(fila, f"No encuentro «{identificador}» con precio en internet. Si no cotiza, "
                                      "créalo en Productos con «Valor anotado a mano» y ese mismo "
                                      "identificador, y vuelve a importar.")
        elif nombre:
            tipo = tipo_prod if tipo_prod in motor.TIPOS else "otro"
            ref = self._nuevo({"nombre": nombre, "tipo": tipo, "fuente": "manual",
                               "largoPlazo": tipo not in ("efectivo", "deuda")}, None)
        else:
            self.plan.error(fila, "No tiene identificador ni nombre: no sé de qué producto es.")
        self.cache[clave] = ref
        return ref

    def _nuevo(self, datos, resultado):
        ref = f"nuevo:{len(self.plan.productos_nuevos) + 1}"
        self.plan.productos_nuevos.append({
            "ref": ref, "datos": datos,
            "precio": resultado and resultado.get("precio"), "monedaPrecio": resultado and resultado.get("moneda"),
            "fecha": resultado and resultado.get("fecha"), "mercado": resultado and resultado.get("mercado")})
        return ref

    def datos(self, ref):
        if ref and ref.startswith("nuevo:"):
            return next(n["datos"] for n in self.plan.productos_nuevos if n["ref"] == ref)
        return almacen.producto(self.cfg, ref)

    def serie(self, ref, fecha=None):
        """Precios en euros del producto; si no llegan a 'fecha', se bajan de nuevo."""
        if ref not in self.series:
            p = self.datos(ref)
            self.series[ref] = motor.serie_producto(p, self.carpeta) if p and p.get("fuente") != "manual" else {}
        s = self.series[ref]
        if fecha and s and max(s) < fecha and ref not in self.frescas:
            self.frescas.add(ref)
            self.series[ref] = s = motor.serie_producto(self.datos(ref), self.carpeta, fresca=True)
        return s


def preparar_tabla(cfg, filas, carpeta):
    """Plantilla o texto de la IA -> Plan."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    cambios = {}
    for n, f in filas:
        fecha = lee_fecha(f.get("fecha"))
        tipo = TIPO_MOV.get(sin_tildes(f.get("tipo_movimiento")))
        ident = str(f.get("identificador") or "").strip()
        nombre = str(f.get("nombre") or "").strip()
        tprod = TIPO_PROD.get(sin_tildes(f.get("tipo_producto")), "")
        importe, unidades, comision = (lee_numero(f.get(k)) for k in ("importe", "unidades", "comision"))
        moneda = (str(f.get("moneda") or "").strip() or "EUR").upper()
        nota = str(f.get("nota") or "").strip()

        if not fecha:
            plan.error(n, f"La fecha «{f.get('fecha') or ''}» no es válida. Usa el formato 2025-03-10 o 10/03/2025.")
            continue
        if not tipo:
            plan.error(n, f"El tipo de movimiento «{f.get('tipo_movimiento') or ''}» no es válido: "
                          "usa compra, venta, dividendo, comision o saldo.")
            continue
        if "error" in (importe, unidades, comision):
            plan.error(n, "Algún número no se entiende (importe, unidades o comisión).")
            continue
        if not importe and tipo != "saldo":
            plan.error(n, "No tiene importe.")
            continue
        if not ident and not nombre:
            plan.error(n, "No tiene identificador: pon el ISIN o el ticker (o, si es una cuenta, su nombre).")
            continue
        if tipo == "saldo" and not ident:
            tprod = tprod or "efectivo"
        ref = res.producto(n, ident, nombre, tprod)
        if not ref:
            continue
        marcas = []

        # Importes en otra moneda: se pasan a euros con el cambio de ese día.
        if moneda != "EUR" and importe is not None:
            if moneda not in cambios:
                cambios[moneda] = motor.serie_cambio(moneda, carpeta)
            fx = valor_en(cambios[moneda], fecha, margen=6)
            if not fx:
                plan.error(n, f"No encuentro el cambio de {moneda} a euros del {almacen.fmt_fecha(fecha)}.")
                continue
            importe = importe * fx
            comision = comision * fx if comision else comision
            marcas.append(f"convertido de {moneda}")

        if tipo == "saldo":
            plan.valoraciones.append({"fila": n, "producto": ref, "fecha": fecha, "valor": round(importe or 0, 2)})
            continue

        p = res.datos(ref)
        cotiza = p and p.get("fuente") != "manual" and p.get("tipo") not in almacen.SOLO_SALDO
        if tipo in ("compra", "venta") and not unidades and cotiza:
            # Muchos extractos no dan las participaciones: se calculan con el precio de ese día.
            precio = valor_en(res.serie(ref, fecha), fecha, margen=6)
            if not precio:
                plan.error(n, "No trae unidades y no encuentro el precio de ese día para calcularlas.")
                continue
            neto = importe - (comision or 0) if tipo == "compra" else importe + (comision or 0)
            unidades = neto / precio
            marcas.append("unidades calculadas")
        plan.movimientos.append({"fila": n, "producto": ref, "fecha": fecha, "tipo": tipo,
                                 "unidades": unidades, "importe": round(importe, 2),
                                 "comision": round(comision, 2) if comision else None,
                                 "nota": nota, "marcas": marcas, "origen": "importado"})
    return plan


def preparar_myinvestor(cfg, archivos, carpeta):
    """[(nombre de archivo, bytes)] de MyInvestor -> Plan. Cada archivo sustituye lo que se
    importó antes de ese fondo, porque el extracto siempre trae la foto completa."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    for nombre, contenido in archivos:
        m = re.search(r"([A-Z]{2}[A-Z0-9]{9}\d)", (nombre or "").upper())
        if not m:
            plan.error(nombre, "No encuentro el ISIN en el nombre del archivo. Descárgalo otra vez de "
                               "MyInvestor sin cambiarle el nombre (lleva el ISIN del fondo).")
            continue
        isin = m.group(1)
        lotes, reembolsos = leer_csv_myinvestor(contenido)
        if not lotes and not reembolsos:
            plan.error(nombre, "No parece un extracto «Plusvalías y minusvalías» de MyInvestor, o está vacío.")
            continue
        ref = res.producto(nombre, isin, "", "fondo")
        if not ref:
            continue
        p = res.datos(ref)
        if p.get("fuente") == "manual":
            plan.error(nombre, f"«{p.get('corto') or p['nombre']}» se valora a mano: el extracto de MyInvestor "
                               "necesita un fondo con precio en internet.")
            continue
        serie = res.serie(ref, lotes[-1][0] if lotes else None)
        movs, sin_vl = myinvestor_a_movimientos(ref, lotes, reembolsos, serie)
        if sin_vl:
            plan.avisos.append(f"{isin}: {sin_vl} compras sin valor liquidativo de ese día; sus "
                               "participaciones se han estimado con el valor del extracto.")
        plan.reemplazar.append((ref, "myinvestor"))
        for mv in movs:
            plan.movimientos.append({**mv, "fila": nombre, "comision": None, "marcas": [],
                                     "origen": "myinvestor"})
    return plan


# ---------------------------------------------------------------- banco (ingresos/gastos)

# Cabeceras que se aceptan en el CSV/extracto del banco.
SINONIMOS_BANCO = {
    "fecha": ["fecha", "date", "fecha operacion", "fecha valor", "fecha contable", "f valor", "f operacion", "fecha de la operacion"],
    "concepto": ["concepto", "descripcion", "description", "detalle", "movimiento", "beneficiario",
                 "nombre", "referencia", "observaciones", "concepto de la operacion"],
    "importe": ["importe", "amount", "cantidad", "importe eur", "importe (eur)", "total"],
    "cargo": ["cargo", "cargos", "debe", "gasto", "salida", "pago", "retiro", "debito"],
    "abono": ["abono", "abonos", "haber", "ingreso", "entrada", "cobro", "credito"],
    "tipo": ["tipo", "tipo movimiento", "tipo de movimiento", "ingreso/gasto"],
    "categoria": ["categoria", "category", "categoría"],
    "titular": ["titular", "titulares", "cuenta", "owner"],
}

# Palabra clave en el concepto -> nombre de categoría. Se aplica solo si esa categoría
# existe en la configuración del usuario (del tipo correcto); si no, se deja sin categoría.
REGLAS_CATEGORIA = {
    "gasto": {
        "Alimentación": ["mercadona", "carrefour", "lidl", "aldi", " dia ", "consum", "eroski", "alcampo",
                         "supermercado", "super ", "fruteria", "panaderia", "carniceria", "hipercor", "ahorramas"],
        "Transporte": ["gasolinera", "repsol", "cepsa", " bp ", "shell", "galp", "renfe", "metro", " emt ",
                       "taxi", "cabify", "uber", "parking", "aparcamiento", "autopista", "peaje", "dgt", " itv ", "bicimad"],
        "Ocio y restaurantes": ["restaurante", "bar ", "cafeteria", "cafe ", "netflix", "spotify", " hbo", "disney",
                                "cine", "amazon prime", "mcdonald", "burger", "telepizza", "dominos", "glovo",
                                "just eat", "uber eats", "steam", "playstation", "gimnasio", "decathlon"],
        "Vivienda": ["alquiler", "hipoteca", "comunidad", "administrador de fincas", " ibi", "seguro hogar"],
        "Suministros": ["iberdrola", "endesa", "naturgy", "movistar", "vodafone", "orange", "masmovil", "yoigo",
                        "digi", "agua", "canal de isabel", "gas natural", "electricidad", "fibra", "internet", "luz "],
        "Salud": ["farmacia", "clinica", "dentista", "hospital", "sanitas", "adeslas", " dkv", "optica", "seguro de salud"],
    },
    "ingreso": {
        "Nómina": ["nomina", "nómina", "salario", "paga", "transferencia nomina", "haberes"],
        "Alquileres cobrados": ["alquiler cobrado", "renta", "arrendamiento"],
        "Intereses y dividendos": ["intereses", "interes", "dividendo", "dividendos", "cupon", "rendimiento"],
    },
}


def leer_tabla_banco(nombre, contenido):
    """Extracto del banco (Excel, CSV o texto pegado) -> [(fila, {columna: valor})].
    Flexible: reconoce fecha + (importe con signo, o columnas cargo/abono) + concepto."""
    if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        hoja = libro["Movimientos"] if "Movimientos" in libro.sheetnames else libro.worksheets[0]
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    else:
        texto = decodifica(contenido)
        texto = re.sub(r"^\s*```[a-zA-Z]*\s*$", "", texto, flags=re.M).strip()
        lineas = [l for l in texto.splitlines() if l.strip()]
        if not lineas:
            return [], "El texto está vacío."
        delim = max([";", "\t", ","], key=lineas[0].count)
        filas = list(csv.reader(io.StringIO("\n".join(lineas)), delimiter=delim))
    filas = [(n, f) for n, f in enumerate(filas, start=1)
             if f and any(c not in (None, "") and str(c).strip() for c in f)]
    if not filas:
        return [], "No hay ninguna fila con datos."
    cab = [sin_tildes(c) for c in filas[0][1]]
    idx = {}
    for col, alias in SINONIMOS_BANCO.items():
        for i, c in enumerate(cab):
            if c in alias and i not in idx.values():
                idx[col] = i
                break
    if "fecha" not in idx:
        return [], ("No encuentro la columna de la fecha. La primera fila tiene que ser la cabecera; "
                    "lo normal en un extracto es: Fecha; Concepto; Importe.")
    if "importe" not in idx and not ("cargo" in idx or "abono" in idx):
        return [], ("No encuentro la columna del importe (ni «cargo»/«abono»). Añade una columna «importe» "
                    "con el signo (negativo = gasto) o deja las de cargo y abono de tu banco.")
    if len(filas) < 2:
        return [], "Solo hay cabecera: debajo deben ir tus movimientos del banco, uno por fila."
    salida = [(n, {col: (f[i] if i < len(f) else None) for col, i in idx.items()}) for n, f in filas[1:]]
    return salida, None


def _importe_banco(f):
    """Devuelve (tipo, importe_positivo) a partir de la fila. tipo = 'ingreso'/'gasto'.
    Usa la columna 'tipo' si existe; si no, el signo del importe, o cargo/abono."""
    imp = lee_numero_con_signo(f.get("importe"))
    tipo_col = sin_tildes(f.get("tipo"))
    if imp == "error":
        return None, "error"
    if imp is None:
        cargo = lee_numero(f.get("cargo"))
        abono = lee_numero(f.get("abono"))
        if "error" in (cargo, abono):
            return None, "error"
        if abono:
            return "ingreso", abs(abono)
        if cargo:
            return "gasto", abs(cargo)
        return None, None
    if tipo_col in ("ingreso", "ingresos", "abono", "haber", "entrada"):
        return "ingreso", abs(imp)
    if tipo_col in ("gasto", "gastos", "cargo", "debe", "salida", "pago"):
        return "gasto", abs(imp)
    # Sin columna de tipo: el signo manda (negativo = gasto).
    return ("gasto" if imp < 0 else "ingreso"), abs(imp)


def lee_numero_con_signo(v):
    """Como lee_numero pero CONSERVANDO el signo (para el importe del banco)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("€", "").replace("$", "").replace(" ", "")
    neg = t.startswith("-") or (t.startswith("(") and t.endswith(")"))
    t = t.lstrip("+-").strip("()")
    if not re.fullmatch(r"[\d.,]+", t):
        return "error"
    n = num_es(t)
    return -n if neg else n


def _categoria_auto(concepto, tipo, cats_del_tipo, reglas_usuario=None):
    """Asigna una categoría por palabra clave, solo si existe en la config del usuario.
    Las reglas del usuario (reglas_usuario = [(palabra, categoria)]) tienen prioridad."""
    disponibles = {sin_tildes(c): c for c in cats_del_tipo}
    conc = " " + sin_tildes(concepto) + " "
    for palabra, categoria in (reglas_usuario or []):
        if sin_tildes(categoria) in disponibles and sin_tildes(palabra) in conc:
            return disponibles[sin_tildes(categoria)]
    for cat, claves in REGLAS_CATEGORIA.get(tipo, {}).items():
        if sin_tildes(cat) in disponibles and any(sin_tildes(k) in conc for k in claves):
            return disponibles[sin_tildes(cat)]
    return ""


# Palabras demasiado genéricas para convertirse en una regla aprendida.
_STOP_PALABRAS = {"compra", "pago", "pagos", "recibo", "recibos", "transferencia", "transf",
                  "tarjeta", "bizum", "cargo", "abono", "ingreso", "domiciliacion", "factura",
                  "nomina", "comision", "comisiones", "liquidacion", "operacion", "cuenta",
                  "madrid", "barcelona", "espana", "online", "web", "www"}


def palabra_clave(concepto):
    """Extrae del concepto una palabra representativa para una regla aprendida
    (la más larga, ≥4 letras, que no sea genérica). '' si no hay ninguna buena."""
    tokens = re.findall(r"[a-záéíóúñ]{4,}", sin_tildes(str(concepto or "")).lower())
    tokens = [t for t in tokens if t not in _STOP_PALABRAS]
    return max(tokens, key=len) if tokens else ""


def aprende_reglas(cfg, plan, cambios_cat):
    """A partir de las categorías que el usuario corrigió en la vista previa del banco,
    guarda reglas propias (palabra del concepto -> categoría) para la próxima vez."""
    conf = cfg.setdefault("config", {})
    reglas = conf.setdefault("reglasCategoria", [])
    existentes = {r["palabra"] for r in reglas}
    por_fila = {str(fl.get("fila")): fl for fl in getattr(plan, "flujos", [])}
    añadidas = 0
    for fila, categoria in (cambios_cat or {}).items():
        categoria = str(categoria or "").strip()
        fl = por_fila.get(str(fila))
        if not categoria or not fl:
            continue
        palabra = palabra_clave(fl.get("concepto"))
        if palabra and palabra not in existentes:
            reglas.append({"palabra": palabra, "categoria": categoria[:40]})
            existentes.add(palabra)
            añadidas += 1
    return añadidas


def preparar_banco(cfg, filas):
    """Extracto del banco -> Plan con flujos (ingresos/gastos)."""
    plan = Plan()
    plan.tipoImport = "banco"
    cats = (cfg.get("config") or {}).get("categorias") or []
    por_tipo = {"ingreso": [c["nombre"] for c in cats if c.get("tipo") == "ingreso"],
                "gasto": [c["nombre"] for c in cats if c.get("tipo") == "gasto"]}
    reglas_usuario = [(r.get("palabra", ""), r.get("categoria", ""))
                      for r in ((cfg.get("config") or {}).get("reglasCategoria") or [])]
    titulares = cfg.get("titulares") or []
    for n, f in filas:
        fecha = lee_fecha(f.get("fecha"))
        if not fecha:
            plan.error(n, f"La fecha «{f.get('fecha') or ''}» no es válida. Usa 2025-03-10 o 10/03/2025.")
            continue
        tipo, importe = _importe_banco(f)
        if importe == "error":
            plan.error(n, "El importe no se entiende.")
            continue
        if not tipo or not importe:
            plan.error(n, "La fila no tiene importe (ni cargo/abono).")
            continue
        concepto = str(f.get("concepto") or "").strip()
        cat_dada = str(f.get("categoria") or "").strip()
        disp = {sin_tildes(c): c for c in por_tipo.get(tipo, [])}
        categoria = disp.get(sin_tildes(cat_dada), "") if cat_dada else _categoria_auto(concepto, tipo, por_tipo.get(tipo, []), reglas_usuario)
        tit = str(f.get("titular") or "").strip()
        titular = next((t for t in titulares if sin_tildes(t) == sin_tildes(tit)), tit if tit else "")
        plan.flujos.append({"fila": n, "fecha": fecha, "tipo": tipo, "importe": round(abs(importe), 2),
                            "categoria": categoria, "titular": titular, "concepto": concepto[:120]})
    return plan


# ---------------------------------------------------------------- aplicar

def aplicar(cfg, plan):
    """Ejecuta el plan sobre cfg (la real o una copia) y devuelve el informe."""
    ids, errores = {}, list(plan.errores)
    for n in plan.productos_nuevos:
        prod, _ = almacen.guarda_producto(cfg, dict(n["datos"]))
        ids[n["ref"]] = prod["id"]
    def real(ref):
        return ids.get(ref, ref)

    sustituidos = 0
    for ref, origen in plan.reemplazar:
        antes = len(cfg["movimientos"])
        cfg["movimientos"] = [m for m in cfg["movimientos"]
                              if not (m.get("producto") == real(ref) and m.get("origen") == origen)]
        sustituidos += antes - len(cfg["movimientos"])

    def firma(m):
        return (m["producto"], m["fecha"], m["tipo"], round(float(m.get("unidades") or 0), 4),
                round(float(m.get("importe") or 0), 2))
    existentes = {firma(m) for m in cfg["movimientos"]}

    añadidos, duplicados, filas = 0, 0, []
    orden = {"compra": 0, "comision": 1, "dividendo": 2, "venta": 3}
    for mv in sorted(plan.movimientos, key=lambda m: (m["fecha"], orden.get(m["tipo"], 9))):
        datos = {k: mv[k] for k in ("fecha", "tipo", "unidades", "importe", "comision", "nota") if mv.get(k) is not None}
        datos["producto"] = real(mv["producto"])
        estado = "nuevo"
        if firma({**datos, "unidades": datos.get("unidades", 0)}) in existentes:
            duplicados += 1
            estado = "repetido"
        else:
            try:
                guardado = almacen.guarda_movimiento(cfg, datos)
                guardado["origen"] = mv["origen"]
                añadidos += 1
            except almacen.ErrorValidacion as e:
                errores.append({"fila": mv["fila"], "mensaje": " ".join(e.errores)})
                estado = "error"
        filas.append({"fila": mv["fila"], "fecha": mv["fecha"], "producto": datos["producto"], "tipo": mv["tipo"],
                      "unidades": mv.get("unidades"), "importe": mv["importe"], "marcas": mv["marcas"],
                      "estado": estado})

    saldos = 0
    for v in plan.valoraciones:
        try:
            almacen.guarda_valoracion(cfg, {"producto": real(v["producto"]), "fecha": v["fecha"], "valor": v["valor"]})
            saldos += 1
            filas.append({"fila": v["fila"], "fecha": v["fecha"], "producto": real(v["producto"]), "tipo": "saldo",
                          "unidades": None, "importe": v["valor"], "marcas": [], "estado": "nuevo"})
        except almacen.ErrorValidacion as e:
            errores.append({"fila": v["fila"], "mensaje": " ".join(e.errores)})

    # Flujos del banco (ingresos y gastos). Dedup por fecha+tipo+importe+concepto.
    def firma_fl(fl):
        return (fl["fecha"], fl["tipo"], round(float(fl.get("importe") or 0), 2),
                (fl.get("nota") or fl.get("concepto") or "")[:60].strip().lower())
    flujos_existentes = {firma_fl(x) for x in cfg.get("flujos", [])}
    flujos_add, flujos_rep, flujos_filas = 0, 0, []
    for fl in plan.flujos:
        datos = {"fecha": fl["fecha"], "tipo": fl["tipo"], "importe": fl["importe"],
                 "categoria": fl.get("categoria") or "", "titular": fl.get("titular") or "",
                 "nota": fl.get("concepto") or ""}
        estado = "nuevo"
        if firma_fl(fl) in flujos_existentes:
            flujos_rep += 1
            estado = "repetido"
        else:
            try:
                almacen.guarda_flujo(cfg, datos)
                flujos_existentes.add(firma_fl(fl))
                flujos_add += 1
            except almacen.ErrorValidacion as e:
                errores.append({"fila": fl["fila"], "mensaje": " ".join(e.errores)})
                estado = "error"
        flujos_filas.append({"fila": fl["fila"], "fecha": fl["fecha"], "tipo": fl["tipo"],
                             "categoria": fl.get("categoria") or "", "concepto": fl.get("concepto") or "",
                             "importe": fl["importe"], "estado": estado})

    def total(t):
        return round(sum(f["importe"] for f in filas if f["tipo"] == t and f["estado"] == "nuevo"), 2)
    nombres = {p["id"]: p.get("corto") or p["nombre"] for p in cfg["productos"]}
    for f in filas:
        f["productoNombre"] = nombres.get(f["producto"], f["producto"])
    def total_fl(t):
        return round(sum(f["importe"] for f in flujos_filas if f["tipo"] == t and f["estado"] == "nuevo"), 2)
    return {
        "tipoImport": plan.tipoImport,
        "añadidos": añadidos, "repetidos": duplicados, "saldos": saldos, "sustituidos": sustituidos,
        "errores": sorted(errores, key=lambda e: (str(type(e["fila"])), str(e["fila"]).zfill(6))),
        "avisos": plan.avisos,
        "totales": {"compras": total("compra"), "ventas": total("venta"),
                    "dividendos": total("dividendo"), "comisiones": total("comision"),
                    "numCompras": sum(1 for f in filas if f["tipo"] == "compra" and f["estado"] == "nuevo")},
        "productosNuevos": [{**n["datos"], "id": ids.get(n["ref"]), "precio": n["precio"],
                             "monedaPrecio": n["monedaPrecio"], "fechaPrecio": n["fecha"],
                             "mercado": n["mercado"]} for n in plan.productos_nuevos],
        "filas": sorted(filas, key=lambda f: f["fecha"], reverse=True),
        "flujosAñadidos": flujos_add, "flujosRepetidos": flujos_rep,
        "flujosFilas": sorted(flujos_filas, key=lambda f: f["fecha"], reverse=True),
        "totalesBanco": {"ingresos": total_fl("ingreso"), "gastos": total_fl("gasto"),
                         "numIngresos": sum(1 for f in flujos_filas if f["tipo"] == "ingreso" and f["estado"] == "nuevo"),
                         "numGastos": sum(1 for f in flujos_filas if f["tipo"] == "gasto" and f["estado"] == "nuevo")},
    }


def vista_previa(cfg, plan):
    """El mismo informe que dará aplicar(), sin tocar la cartera de verdad."""
    return aplicar(copy.deepcopy(cfg), plan)
