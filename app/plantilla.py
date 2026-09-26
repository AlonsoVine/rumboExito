# -*- coding: utf-8 -*-
"""
plantilla.py  ·  La plantilla de importación, en Excel y en CSV
===============================================================
"""

import datetime as dt
import io

from .importar import COLUMNAS

TIPOS_PRODUCTO = ["fondo", "etf", "accion", "cripto", "commodity", "bono", "pension",
                  "efectivo", "inmueble", "deuda", "otro"]
TIPOS_MOVIMIENTO = ["compra", "venta", "dividendo", "comision", "saldo"]

EXPLICACION = [
    ("fecha", "Sí", "Día de la operación. Por ejemplo 10/03/2025."),
    ("identificador", "Casi siempre", "ISIN (IE00BYX5NX33) o ticker (AAPL, BTC-EUR). Vacío solo en cuentas y cosas sin precio en internet."),
    ("nombre", "Si no hay identificador", "Nombre del producto o de la cuenta. Sirve para reconocer las cuentas bancarias."),
    ("tipo_producto", "No", "fondo, etf, accion, cripto, commodity, bono, pension, efectivo, inmueble, deuda u otro."),
    ("tipo_movimiento", "Sí", "compra, venta, dividendo, comision o saldo (el saldo de una cuenta en esa fecha)."),
    ("unidades", "Recomendado", "Participaciones, acciones u onzas. Si lo dejas vacío, se calculan con el precio de ese día."),
    ("importe", "Sí", "Dinero total. Compra: lo que salió de tu cuenta, con comisiones. Venta o dividendo: lo que entró."),
    ("moneda", "No", "Moneda del importe. Si no es EUR, se convierte con el cambio de ese día. Vacío = EUR."),
    ("comision", "No", "Comisión de la operación, ya incluida en el importe."),
    ("nota", "No", "Lo que quieras apuntar."),
]
EJEMPLO = [
    (dt.date(2025, 3, 10), "IE00BYX5NX33", "Fidelity MSCI World", "fondo", "compra", 62.5, 900, "EUR", None, "Aportación mensual"),
    (dt.date(2025, 4, 2), "AAPL", "Apple", "accion", "compra", 5, 1050, "USD", 2, ""),
    (dt.date(2025, 5, 15), "BTC-EUR", "Bitcoin", "cripto", "compra", 0.01, 600, "EUR", None, ""),
    (dt.date(2025, 6, 30), "AAPL", "Apple", "accion", "dividendo", None, 1.25, "USD", None, ""),
    (dt.date(2025, 6, 30), "", "Cuenta del banco", "efectivo", "saldo", None, 2500, "EUR", None, "Saldo a fin de mes"),
]


def csv_vacio():
    return (";".join(COLUMNAS) + "\n").encode("utf-8-sig")


def excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    libro = Workbook()
    cab_estilo = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="2A78D6"),
                      alignment=Alignment(vertical="center"))
    anchos = [13, 16, 28, 15, 17, 12, 12, 9, 11, 30]

    def prepara(hoja, filas):
        hoja.append(COLUMNAS)
        for c in hoja[1]:
            for k, v in cab_estilo.items():
                setattr(c, k, v)
        for i, a in enumerate(anchos):
            hoja.column_dimensions[chr(65 + i)].width = a
        hoja.row_dimensions[1].height = 22
        hoja.freeze_panes = "A2"
        for f in filas:
            hoja.append(list(f))
        for fila in hoja.iter_rows(min_row=2, max_row=max(hoja.max_row, 500)):
            fila[0].number_format = "DD/MM/YYYY"
            for c in fila[5:7] + fila[8:9]:
                c.number_format = "#,##0.00######"

    hoja = libro.active
    hoja.title = "Movimientos"
    prepara(hoja, [])
    for col, lista in (("D", TIPOS_PRODUCTO), ("E", TIPOS_MOVIMIENTO)):
        dv = DataValidation(type="list", formula1='"' + ",".join(lista) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor no válido",
                            error="Elige una opción de la lista.")
        dv.add(f"{col}2:{col}2000")
        hoja.add_data_validation(dv)

    ej = libro.create_sheet("Ejemplo")
    prepara(ej, EJEMPLO)

    ins = libro.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 22
    ins.column_dimensions["C"].width = 100
    ins.append(["Cómo rellenar la hoja «Movimientos»"])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append(["Una fila por operación. Mira la hoja «Ejemplo». Cuando termines, guarda el archivo e impórtalo "
                "desde la app: Mis datos → Importar → Plantilla."])
    ins.append([])
    ins.append(["Columna", "¿Obligatoria?", "Qué poner"])
    for c in ins[4]:
        for k, v in cab_estilo.items():
            setattr(c, k, v)
    for fila in EXPLICACION:
        ins.append(list(fila))
    for fila in ins.iter_rows(min_row=5):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
